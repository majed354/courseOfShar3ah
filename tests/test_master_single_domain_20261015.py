import copy
import hashlib
import json
from pathlib import Path
import re
import unittest
import pymupdf
from docx import Document
from scripts.verify_course_outcomes import ErrorCollector,validate_single_domain_policy
from _shared_alignment_test_support import lineage_records, previous_context

ROOT=Path(__file__).resolve().parents[1]
B=ROOT/'assets/course-specifications/quranic-master-single-domain-20261015'
def read(p):return json.loads(p.read_text())
def effective(v):
    rows=copy.deepcopy(v['extracted']['clos'])
    for o in v['overrides']:
        if o['field']=='clos':rows=o['value']
    return rows
class MasterSingleDomainTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg=read(B/'revision-decisions.json');cls.out=read(ROOT/'course-outcomes.json');cls.data=read(ROOT/'data.json');cls.ref=read(B/'program-reference.json');cls.before=read(B/'previous-context.json')
    def test_94_real_coded_outcomes_one_link_in_own_domain(self):
        e=ErrorCollector();validate_single_domain_policy(self.out,e);self.assertEqual([],e.messages)
        n=0
        for c in self.cfg['courses']:
            v=next(v for v in self.out['courses'][c['code']]['variants'] if '/quranic-master-single-domain-' in v['source_pdf'])
            rows=effective(v);self.assertEqual(len(c['clos']),len(rows))
            for x,y in zip(rows,c['clos']):
                self.assertEqual((y['code'],y['text'],[y['plo_code']]),(x['code'],x['text'],x['plo_mappings'][0]['plo_codes']))
                self.assertEqual(x['document_plo_codes'],[y['plo_code']]);self.assertTrue(x['assessment']);n+=1
        self.assertEqual(94,n)
    def test_policy_rejects_multiple_cross_domain_and_uncoded_rows(self):
        for field,value in [('code',None),('plo_mappings',[{'scope':{'program':self.cfg['program']},'plo_codes':['ع1','ع2']}]),('plo_mappings',[{'scope':{'program':self.cfg['program']},'plo_codes':['ق1']}])]:
            row={'code':'1.1','text':'اختبار','plo_mappings':[{'scope':{'program':self.cfg['program']},'plo_codes':['ع1']}]};row[field]=value
            p={'reference_policy':{'single_domain_programs':[self.cfg['program']]},'courses':{'x':{'variants':[{'scopes':[{'program':self.cfg['program']}],'extracted':{'clos':[row]}}]}}}
            e=ErrorCollector();validate_single_domain_policy(p,e);self.assertGreater(e.count,0,field)
    def test_plan_plo_texts_and_other_contexts_preserved(self):
        p=next(p for p in self.data['programs'] if p['name']==self.cfg['program']);old=self.before['program']
        self.assertEqual(old['courses'],p['courses']);self.assertEqual((14,34),(len(p['courses']),sum(c['hours'] for c in p['courses'])))
        self.assertEqual(old['program_details']['learning_outcomes'],p['program_details']['learning_outcomes'])
        import subprocess
        baseline=json.loads(subprocess.check_output(['git','show',self.before['baseline_commit']+':course-outcomes.json'],cwd=ROOT))
        doctorate=read(ROOT/'assets/course-specifications/quranic-doctorate-single-domain-20261016/previous-context.json')['course_outcomes']
        _,before_specialization=previous_context()
        specialized_codes={r['course_code'] for r in lineage_records()}
        for code,course in baseline['courses'].items():
            if code in doctorate:self.assertEqual(course,doctorate[code])
            elif code not in self.before['course_outcomes']:
                # Preserve this earlier batch's outside-context comparison for
                # the later specialized identities; live descendants have
                # their own complete source/content/scope contract.
                current=before_specialization['courses'][code] if code in specialized_codes else self.out['courses'][code]
                self.assertEqual(course,current,code)
            else:
                old=[v for v in course['variants'] if not any(s['program']==self.cfg['program'] for s in v['scopes'])]
                new=[v for v in self.out['courses'][code]['variants'] if not any(s['program']==self.cfg['program'] for s in v['scopes'])]
                self.assertEqual(old,new)
    def test_69_supported_matrix_links_and_17_mastery_links(self):
        matrix=self.ref['program_matrix'];self.assertEqual(69,len(matrix));self.assertEqual(17,sum(x['classification']==3 for x in matrix))
        for c in self.cfg['courses']:
            actual={r['plo_code']:r['classification'] for r in matrix if r['course_code']==c['code']}
            self.assertEqual(c['matrix'],actual);self.assertEqual({x['plo_code'] for x in c['clos']},set(actual))
        for plo,mastery in self.cfg['mastery'].items():self.assertEqual(set(mastery),{r['course_code'] for r in matrix if r['plo_code']==plo and r['classification']==3})
    def test_q2_has_own_values_outcome_and_evidence_not_skill_grade(self):
        project=next(c for c in self.cfg['courses'] if c['code']=='2002888-4');rows={x['code']:x for x in project['clos']}
        self.assertEqual('ق2',rows['3.2']['plo_code']);self.assertEqual('م2',rows['2.2']['plo_code']);self.assertEqual('ق1',rows['3.1']['plo_code'])
        rubric=read(B/'q2-project-rubric.json');self.assertEqual('3.2',rubric['clo_code']);self.assertEqual(['2.2','3.1'],rubric['separate_from']);self.assertEqual(100,sum(x['weight'] for x in rubric['criteria']));self.assertEqual(4,len(rubric['criteria']));self.assertTrue(all(x['evidence'] for x in rubric['criteria']))
        self.assertFalse(read(B/'manifest.json')['historical_grades_changed']);self.assertFalse(read(B/'manifest.json')['institutional_approval_claimed'])
    def test_source_hashes_page_locators_word_tables_and_unchanged_weights(self):
        for file in read(B/'manifest.json')['files']:self.assertEqual(file['sha256'],hashlib.sha256((ROOT/file['path']).read_bytes()).hexdigest(),file['path'])
        pages=read(B/'clo-page-map.json');self.assertEqual(94,len(pages))
        for c in self.cfg['courses']:
            word=Document(B/'courses'/c['code']/'course-specification.docx')
            table=word.tables[2];self.assertEqual([[x['code'],x['text'],x['plo_code'],x['teaching'],x['assessment']] for x in c['clos']],[[a.text for a in row.cells] for row in table.rows[1:]])
            self.assertEqual(['',''],[word.tables[1].rows[1].cells[i].text for i in (1,2)])
            self.assertEqual(100,sum(x['weight'] for x in c['assessment_plan_preserved']))
            source=ROOT/c['pdf'];self.assertEqual(c['sha256'],hashlib.sha256(source.read_bytes()).hexdigest())
        for loc in pages:
            doc=pymupdf.open(ROOT/loc['source_pdf']);box=pymupdf.Rect(loc['code_bbox']);self.assertIn(loc['clo_code'],doc[loc['source_page']-1].get_textbox(box+(-1,-1,1,1)))
    def test_program_word_and_course_links_match_current_matrix(self):
        word=Document(B/'program-specification.docx')
        for row,c in zip(word.tables[9].rows[3:],self.cfg['courses']):self.assertEqual([str(c['matrix'].get(o['code'],'')) for o in self.cfg['outcomes']],[x.text for x in row.cells[1:]])
        self.assertEqual(95,len(word.tables[31].rows));self.assertEqual('',word.tables[30].rows[1].cells[1].text);self.assertEqual('',word.tables[30].rows[2].cells[1].text)
        links=[r.target_ref for r in word.part.rels.values() if 'hyperlink' in r.reltype and '/courses/' in r.target_ref];self.assertEqual(14,len(links));self.assertTrue(all('quranic-master-single-domain-' in x for x in links))
