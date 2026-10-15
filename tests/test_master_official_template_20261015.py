import hashlib
import json
from pathlib import Path
import unittest
import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / 'assets/course-specifications/quranic-master-official-template-20261015'
PREVIOUS = ROOT / 'assets/course-specifications/quranic-master-plo-update-20261015'
def read(path): return json.loads(path.read_text(encoding='utf-8'))

class OfficialTemplatePublicationTests(unittest.TestCase):
    def test_sources_are_exact_user_provided_bytes(self):
        manifest=read(BUNDLE/'manifest.json')
        self.assertTrue(manifest['source_bytes_preserved'])
        self.assertEqual('8147b9d383066ed9f64da6b9c50ef2892c2a4928cdd8d6c4e67a2c5dce44c2c3',hashlib.sha256((BUNDLE/'program-specification.pdf').read_bytes()).hexdigest())
        self.assertEqual('cffdf3d38c42975a7f7bddb888bdced69d9c2036293d1533a23d5428e9073439',hashlib.sha256((BUNDLE/'program-specification.docx').read_bytes()).hexdigest())
        for file in manifest['files']:
            self.assertEqual(file['sha256'],hashlib.sha256((ROOT/file['path']).read_bytes()).hexdigest())
        self.assertEqual(22,len(pymupdf.open(BUNDLE/'program-specification.pdf')))

    def test_only_source_locations_change_in_the_program_reference(self):
        before=read(PREVIOUS/'program-reference.json');after=read(BUNDLE/'program-reference.json')
        for key in ['program_outcomes','curriculum','program_matrix']:
            trim=lambda rows:[{k:v for k,v in x.items() if k not in {'source_page','source_file'}} for x in rows]
            self.assertEqual(trim(before[key]),trim(after[key]))
        self.assertEqual({5},{p['source_page'] for p in after['program_outcomes']})
        self.assertEqual({6},{p['source_page'] for p in after['curriculum']})
        self.assertEqual({7,8},{p['source_page'] for p in after['program_matrix']})
        self.assertEqual(hashlib.sha256((BUNDLE/'program-specification.pdf').read_bytes()).hexdigest(),after['program_metadata']['Source_Matrix_Review_SHA256'])

    def test_all_scoped_clo_locators_use_the_new_pdf_and_its_exact_page(self):
        page_map={(x['course_code'],x['clo_code']):x['source_page'] for x in read(BUNDLE/'clo-page-map.json')}
        self.assertEqual(93,len(page_map))
        data={'courses':read(ROOT/'assets/course-specifications/quranic-master-single-domain-20261015/previous-context.json')['course_outcomes']};digest=read(BUNDLE/'manifest.json')['source_files']['pdf']['sha256'];count=0
        for code,course in data['courses'].items():
            for variant in course['variants']:
                for override in variant['overrides']:
                    if override.get('policy')!='authorized_program_alignment_revision':continue
                    self.assertIn(digest,override['evidence'])
                    index=int(override['field'].split('[')[1].split(']')[0]);clo=variant['extracted']['clos'][index]['code'] or 'غير مرمز'
                    for mapping in override['value']:
                        if digest in mapping['evidence']:
                            self.assertEqual(page_map[(code,clo)],mapping['source_page'])
                            self.assertIn('quranic-master-official-template-20261015',mapping['evidence'])
                    count+=1
        self.assertEqual(93,count)

    def test_published_selection_uses_new_edition_and_keeps_the_curriculum(self):
        program=read(ROOT/'assets/course-specifications/quranic-master-single-domain-20261015/previous-context.json')['program']
        prior=read(PREVIOUS/'previous-context.json')['program']
        self.assertEqual(prior['courses'],program['courses'])
        self.assertEqual((14,34),(len(program['courses']),sum(c['hours'] for c in program['courses'])))
        for key in ['program_specification_url','program_specification_docx_url','program_reference_url']:
            self.assertIn('quranic-master-official-template-20261015',program[key])
            self.assertTrue((ROOT/program[key]).is_file())
        self.assertFalse(read(BUNDLE/'manifest.json')['institutional_approval_claimed'])
        doc=Document(BUNDLE/'program-specification.docx')
        self.assertEqual('',doc.tables[30].rows[1].cells[1].text)
        self.assertEqual('',doc.tables[30].rows[2].cells[1].text)
