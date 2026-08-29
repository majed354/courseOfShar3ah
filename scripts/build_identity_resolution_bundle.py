#!/usr/bin/env python3
"""Build the approved course-identity resolution bundle.

Every source remains untouched.  The bundle contains only cases whose target
identity is supported by the plan catalogue and whose reuse/adaptation the user
approved.  Program-specific variants are deliberately emitted as separate
PDFs even when they share a course code.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

import pymupdf
import arabic_reshaper
from bidi.algorithm import get_display
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_RELATIVE = Path("assets/course-specifications/identity-resolutions-20260829")
WORKING_RELATIVE = Path("tmp/pdfs/identity-resolutions-20260829")
COMBINED_ISLAMIC_STUDIES = Path("توصيفات برنامج الدراسات الإسلامية.pdf")
QIRAAT_STAMP = Path("الأختام/photo_1448-03-06 21.11.19.jpeg")
SOFFICE = Path(shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice")
ARIAL = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
_SAKKAL_CACHE = Path.home() / "Library/Group Containers/UBF8T346G9.Office/FontCache"
_SAKKAL_CANDIDATES = sorted(
    _SAKKAL_CACHE.glob("*/CloudFonts/Sakkal Majalla/34866040251.ttf")
)
SAKKAL_BOLD = _SAKKAL_CANDIDATES[0] if _SAKKAL_CANDIDATES else Path("34866040251.ttf")


SOURCE_HASHES = {
    COMBINED_ISLAMIC_STUDIES:
        "979de4cafe38f11ca6f4e98fc2931c9518d3603de9fd91bc81f37e95d5e28f1e",
    Path("assets/course-specifications/islamic-studies-bundle-1444/20024103-2.pdf"):
        "5dacaccdddb0cc9aea32645e8cebb603e57875412944a106b71c5bcb436b50d4",
    Path("assets/course-specifications/quran-old-v39/2002427-2.pdf"):
        "98e94d90d4afbfe5eb3022e17ded34c1238f0dbc4e5b98d1d51465861fe01de0",
    Path("assets/course-specifications/systems-source-2022/20034108-2.pdf"):
        "db99460211b289cfa65152e148fdc60fa2376f895f6f4dc1948b52c250ab81f5",
    Path("assets/course-specifications/quran-old-v39/2002454-2.pdf"):
        "feb0f3fa455329a398a169b184284c323b012aa162fd6c62091cc04b5a9d7d65",
    Path("assets/course-specifications/qiraat-source-1445/2002454-2.pdf"):
        "2b98525a1bdf34cd3936ca04c2578375927dff498cad4241af4d6581bb2257fa",
    Path(
        "ااااالملفات الخام/توصيف المقررات/"
        "توصيفات دكتوراة القراءات/"
        "دراسات في طبقات القراء وأسانيدهم.docx"
    ): "f43960152841978e6584653f7a61ba8bf1b7a5f5ee77113bc9db907258646da0",
}


def run(command: list[str], *, capture: bool = False) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        command,
        check=True,
        text=True,
        stdout=subprocess.PIPE if capture else None,
        stderr=subprocess.PIPE if capture else None,
    )


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def import_helpers() -> dict[str, Any]:
    sys.path.insert(0, str(ROOT / "scripts"))
    from build_raw_recovery_bundle import (  # noqa: PLC0415
        complete_pdf_approval,
        insert_arabic_html,
        most_common_background,
        remove_nonpublic_link_annotations,
        replace_cover_value,
        replace_docx_label,
    )
    from build_specialty_recovery_bundle import add_accessible_identity  # noqa: PLC0415

    return {
        "complete_pdf_approval": complete_pdf_approval,
        "insert_arabic_html": insert_arabic_html,
        "most_common_background": most_common_background,
        "remove_nonpublic_link_annotations": remove_nonpublic_link_annotations,
        "replace_cover_value": replace_cover_value,
        "replace_docx_label": replace_docx_label,
        "add_accessible_identity": add_accessible_identity,
    }


def assert_sources(root: Path) -> None:
    for relative, expected in SOURCE_HASHES.items():
        source = root / relative
        if not source.is_file():
            raise FileNotFoundError(source)
        actual = sha256(source)
        if actual != expected:
            raise RuntimeError(f"Source hash changed: {relative} ({actual})")
    if not (root / QIRAAT_STAMP).is_file():
        raise FileNotFoundError(root / QIRAAT_STAMP)
    for dependency in (SOFFICE, ARIAL, SAKKAL_BOLD):
        if not dependency.is_file():
            raise FileNotFoundError(f"Required local build dependency: {dependency}")


def extract_pages(source: Path, page_range: str, output: Path) -> None:
    run(["qpdf", "--empty", "--pages", str(source), page_range, "--", str(output)])


def replace_rect(
    page: pymupdf.Page,
    rect: pymupdf.Rect,
    value: str,
    helpers: dict[str, Any],
    *,
    font_size: float = 13.0,
    align: str = "right",
) -> dict[str, Any]:
    background = helpers["most_common_background"](page, rect)
    page.add_redact_annot(rect, fill=background, cross_out=False)
    page.apply_redactions(
        images=pymupdf.PDF_REDACT_IMAGE_NONE,
        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
        text=pymupdf.PDF_REDACT_TEXT_REMOVE,
    )
    scale = helpers["insert_arabic_html"](
        page,
        rect,
        value,
        align=align,
        font_size=font_size,
    )
    return {
        "rect_top_points": [round(number, 3) for number in rect],
        "value": value,
        "htmlbox_scale": scale,
    }


def actual_text_hex(text: str) -> str:
    return "FEFF" + text.encode("utf-16-be").hex().upper()


def add_accessible_actual_text(
    pdf: Path,
    *,
    page_index: int,
    text: str,
    working: Path,
    label: str,
) -> None:
    """Add invisible Unicode /ActualText without changing visible layout."""
    reader = PdfReader(str(pdf), strict=False)
    writer = PdfWriter(clone_from=reader)
    page = writer.pages[page_index]
    resources = page.get("/Resources") or DictionaryObject()
    resources = resources.get_object() if hasattr(resources, "get_object") else resources
    fonts = resources.get("/Font") or DictionaryObject()
    fonts = fonts.get_object() if hasattr(fonts, "get_object") else fonts
    fonts[NameObject("/FActualResolution")] = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources
    stream = DecodedStreamObject()
    stream.set_data(
        (
            f"/Span << /ActualText <{actual_text_hex(text)}> >> BDC\n"
            "BT /FActualResolution 1 Tf 3 Tr 36 20 Td (x) Tj ET\nEMC\n"
        ).encode("ascii")
    )
    stream_reference = writer._add_object(stream)
    contents = page.get("/Contents")
    resolved = contents.get_object() if hasattr(contents, "get_object") else contents
    if contents is None:
        page[NameObject("/Contents")] = stream_reference
    elif isinstance(resolved, ArrayObject):
        page[NameObject("/Contents")] = ArrayObject([*resolved, stream_reference])
    else:
        page[NameObject("/Contents")] = ArrayObject([contents, stream_reference])
    temporary = working / f"{pdf.stem}-{label}-accessible.pdf"
    writer.write(str(temporary))
    temporary.replace(pdf)


def edit_cover_pdf(
    source: Path,
    output: Path,
    helpers: dict[str, Any],
    *,
    title: str,
    code: str,
) -> list[dict[str, Any]]:
    document = pymupdf.open(source)
    page = document[0]
    title_edit = helpers["replace_cover_value"](page, "title", title)
    title_edit.update({"kind": "replace_cover_title", "to": title})
    code_edit = helpers["replace_cover_value"](page, "code", code)
    code_edit.update({"kind": "replace_course_code", "to": code})
    document.save(output, garbage=4, deflate=True, clean=True)
    document.close()
    return [title_edit, code_edit]


def edit_systems_cover(
    source: Path,
    output: Path,
    helpers: dict[str, Any],
    *,
    title: str,
    code: str,
) -> list[dict[str, Any]]:
    """Replace values inside the value column without touching table borders."""
    document = pymupdf.open(source)
    page = document[0]
    title_rect = pymupdf.Rect(52.0, 475.65, 398.65, 500.95)
    code_rect = pymupdf.Rect(52.0, 501.55, 398.65, 526.85)
    title_edit = replace_rect(
        page,
        title_rect,
        title,
        helpers,
        font_size=12.0,
        align="center",
    )
    title_edit.update({"kind": "replace_cover_title", "to": title})
    code_edit = replace_rect(
        page,
        code_rect,
        code,
        helpers,
        font_size=12.0,
        align="center",
    )
    code_edit.update({"kind": "replace_course_code", "to": code})
    # Redaction can consume the hairline separating the value and label
    # columns.  Restore the exact source geometry after inserting the text.
    separator = pymupdf.Rect(398.81, 475.556, 399.29, 501.02)
    page.draw_rect(
        separator,
        color=(0, 0, 0),
        fill=(0, 0, 0),
        width=0,
        overlay=True,
    )
    document.save(output, garbage=4, deflate=True, clean=True)
    document.close()
    return [
        title_edit,
        code_edit,
        {
            "kind": "restore_cover_column_separator",
            "rect_top_points": [round(number, 3) for number in separator],
        },
    ]


def edit_legacy_quran_cover(
    source: Path,
    output: Path,
    helpers: dict[str, Any],
    *,
    title: str,
    code: str,
    working: Path,
) -> list[dict[str, Any]]:
    """Place an overlaid replacement above the legacy form XObject."""
    # Remove the old title and code text-show operations before placing the
    # new visible identity.  A visual overlay alone would leave the legacy
    # values in copy/search text underneath the new ones.
    reader = PdfReader(str(source), strict=False)
    writer = PdfWriter(clone_from=reader)
    page = writer.pages[0]
    contents = page.raw_get("/Contents")
    references = list(contents) if isinstance(contents, ArrayObject) else [contents]
    old_title_matrices = (
        b"1 0 0 1 434.35 631.3 Tm",
        b"1 0 0 1 392.23 631.3 Tm",
        b"1 0 0 1 387.07 631.3 Tm",
        b"1 0 0 1 379.51 631.3 Tm",
        b"1 0 0 1 374.35 631.3 Tm",
    )
    old_code_matrices = (
        b"1 0 0 1 429.07 601.18 Tm",
        b"1 0 0 1 424.75 601.18 Tm",
        b"1 0 0 1 402.07 601.18 Tm",
        b"1 0 0 1 371.83 601.18 Tm",
        b"1 0 0 1 364.51 601.18 Tm",
    )
    old_identity_matrices = (*old_title_matrices, *old_code_matrices)
    removed = 0
    for reference in references:
        stream = reference.get_object()
        data = stream.get_data()
        for matrix in old_identity_matrices:
            marker = matrix + b"\r\n"
            if marker not in data:
                continue
            if data.count(marker) != 1:
                raise RuntimeError(f"Unsafe legacy-title matrix count in {source}")
            start = data.index(marker)
            end = data.find(b"\r\nET", start)
            if end < 0 or end - start > 220:
                raise RuntimeError(f"Unsafe legacy-title text block in {source}")
            block = data[start:end + 4]
            if b"TJ" not in block and b"Tj" not in block:
                raise RuntimeError(f"Legacy-title block has no text-show operator in {source}")
            data = data[:start] + matrix + b"\r\nET" + data[end + 4:]
            removed += 1
        stream.set_data(data)
    if removed != len(old_identity_matrices):
        raise RuntimeError(
            f"Expected {len(old_identity_matrices)} legacy identity blocks in "
            f"{source}, found {removed}"
        )
    clean_source = working / f"{output.stem}-legacy-title-removed.pdf"
    writer.write(str(clean_source))

    reader = PdfReader(str(clean_source), strict=False)
    page_width = float(reader.pages[0].mediabox.width)
    page_height = float(reader.pages[0].mediabox.height)
    overlay_path = working / f"{output.stem}-legacy-cover-overlay.pdf"
    # The source uses very thin purple rules around each row.  Restrict the
    # overlays to the exact interiors so those rules remain fully visible.
    title_rect = pymupdf.Rect(315.0, 194.66, 436.5, 223.57)
    code_rect = pymupdf.Rect(315.0, 224.78, 436.5, 253.69)
    title_background = (0.9489127994, 0.9490348697, 0.9488822818)
    code_background = (0.8509193659, 0.8510414362, 0.8508888483)

    title_font = "SakkalMajallaBoldIdentityResolution"
    code_font = "ArialIdentityResolution"
    if title_font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(title_font, str(SAKKAL_BOLD)))
    if code_font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(code_font, str(ARIAL)))
    overlay = canvas.Canvas(
        str(overlay_path), pagesize=(page_width, page_height), pageCompression=1
    )
    for rect, background in (
        (title_rect, title_background),
        (code_rect, code_background),
    ):
        overlay.setFillColorRGB(*background)
        overlay.rect(
            rect.x0,
            page_height - rect.y1,
            rect.width,
            rect.height,
            stroke=0,
            fill=1,
        )
    overlay.setFillColorRGB(0, 0, 0)
    overlay.setFont(title_font, 18.0)
    overlay.drawRightString(
        434.5,
        page_height - 212.4,
        get_display(arabic_reshaper.reshape(title)),
    )
    overlay.setFont(code_font, 14.0)
    overlay.drawRightString(434.5, page_height - 243.5, code)
    # Restore the purple rule between the title and code rows.  The legacy
    # form stores that rule below the XObject over which the replacement sits.
    overlay.setFillColorRGB(0.298, 0.239, 0.557)
    for top_y0, top_y1 in ((223.58, 223.82), (224.54, 224.78)):
        overlay.rect(
            315.0,
            page_height - top_y1,
            121.5,
            top_y1 - top_y0,
            stroke=0,
            fill=1,
        )
    overlay.showPage()
    overlay.save()

    overlay_reader = PdfReader(str(overlay_path), strict=False)
    writer = PdfWriter(clone_from=reader)
    writer.pages[0].merge_page(overlay_reader.pages[0], over=True, expand=False)
    intermediate = working / f"{output.stem}-legacy-cover-edited.pdf"
    writer.write(str(intermediate))
    run(
        [
            "qpdf",
            "--object-streams=generate",
            "--stream-data=compress",
            "--compression-level=9",
            str(intermediate),
            str(output),
        ]
    )
    description = (
        "يركز مقرر تفسير جزء عم على دراسة الجزء الأخير (الثلاثون) من القرآن الكريم، "
        "المعروف بجزء عم؛ ويتناول تفسير الألفاظ الغريبة وفهم المعنى العام للآيات. "
        "كما يشمل المقرر دراسة أسباب النزول، والمرويات المتعلقة بالآيات، والترجيح بين "
        "الأقوال المختلفة. ويهتم أيضًا بالجمع بين الآيات التي قد تبدو متعارضة، وتطبيق "
        "قواعد التفسير، واستنباط الفوائد من الآيات."
    )
    description_rect = pymupdf.Rect(49.5, 296.0, 545.0, 354.0)
    document = pymupdf.open(output)
    description_edit = replace_rect(
        document[1],
        description_rect,
        description,
        helpers,
        font_size=12.8,
        align="right",
    )
    description_edit.update(
        {
            "kind": "replace_course_title_reference",
            "page_1_based": 2,
            "from": "التفسير (8)",
            "to": title,
        }
    )
    description_output = working / f"{output.stem}-description-edited.pdf"
    document.save(description_output, garbage=4, deflate=True, clean=True)
    document.close()
    description_output.replace(output)
    add_accessible_actual_text(
        output,
        page_index=1,
        text=description,
        working=working,
        label="description",
    )
    return [
        {
            "kind": "remove_hidden_legacy_cover_identity",
            "title_text_blocks": len(old_title_matrices),
            "code_text_blocks": len(old_code_matrices),
            "from_title": "التفسير (8)",
            "from_code": "2002427-2",
        },
        {
            "kind": "replace_cover_title_overlay",
            "to": title,
            "rect_top_points": [round(number, 3) for number in title_rect],
        },
        {
            "kind": "replace_course_code_overlay",
            "to": code,
            "rect_top_points": [round(number, 3) for number in code_rect],
        },
        {
            "kind": "restore_cover_row_separator",
            "rects_top_points": [
                [315.0, 223.58, 436.5, 223.82],
                [315.0, 224.54, 436.5, 224.78],
            ],
        },
        description_edit,
    ]


def edit_sira_variant(
    raw_source: Path,
    output: Path,
    helpers: dict[str, Any],
    *,
    program: str,
    approval: dict[str, list[str]],
    working: Path,
) -> list[dict[str, Any]]:
    extracted = working / f"{output.stem}-raw.pdf"
    extract_pages(raw_source, "469-477", extracted)
    document = pymupdf.open(extracted)
    page = document[0]
    changes = [
        {
            "kind": "replace_course_code",
            **replace_rect(
                page,
                pymupdf.Rect(390.5, 439.5, 452.5, 460.5),
                "2001106-2",
                helpers,
                font_size=12.5,
            ),
        },
        {
            "kind": "replace_program",
            **replace_rect(
                page,
                pymupdf.Rect(220.0, 469.5, 458.0, 491.0),
                program,
                helpers,
                font_size=12.5,
            ),
        },
        {
            "kind": "replace_department",
            **replace_rect(
                page,
                pymupdf.Rect(220.0, 499.5, 436.5, 521.0),
                "قسم القراءات",
                helpers,
                font_size=12.5,
            ),
        },
    ]
    total_hours_rect = pymupdf.Rect(158.2, 270.65, 241.6, 293.1)
    total_hours_scale = helpers["insert_arabic_html"](
        document[3],
        total_hours_rect,
        "30",
        align="center",
        font_size=14.0,
    )
    changes.append(
        {
            "kind": "complete_total_learning_hours",
            "page_1_based": 4,
            "value": 30,
            "rect_top_points": [round(number, 3) for number in total_hours_rect],
            "htmlbox_scale": total_hours_scale,
        }
    )
    intermediate = working / f"{output.stem}-cover.pdf"
    document.save(intermediate, garbage=4, deflate=True, clean=True)
    document.close()
    shutil.copy2(intermediate, output)
    approval_cells = helpers["complete_pdf_approval"](output, 9, approval)
    changes.append(
        {
            "kind": "complete_approval_table",
            "page_1_based": 9,
            "values": approval,
            "cells": approval_cells,
        }
    )
    return [{"kind": "extract_pages", "source_pages_inclusive": "469-477"}, *changes]


def replace_hours_glyph(
    source: Path,
    output: Path,
    *,
    page_index: int,
    matrix: bytes,
) -> list[dict[str, Any]]:
    reader = PdfReader(str(source), strict=False)
    writer = PdfWriter(clone_from=reader)
    page = writer.pages[page_index]
    contents = page.raw_get("/Contents")
    references = list(contents) if isinstance(contents, ArrayObject) else [contents]
    needle = matrix + b"\r\n1 g\r\n1 G\r\n[<0572>] TJ"
    replacement = matrix + b"\r\n1 g\r\n1 G\r\n[<0570>] TJ"
    matches = 0
    for reference in references:
        stream = reference.get_object()
        data = stream.get_data()
        count = data.count(needle)
        if not count:
            continue
        if count != 1:
            raise RuntimeError(f"Unsafe hours-glyph match count in {source}: {count}")
        updated = data.replace(needle, replacement, 1)
        stream.set_data(updated)
        matches += 1
    if matches != 1:
        raise RuntimeError(f"Expected one targeted hours glyph in {source}, found {matches}")

    cmap_fixed = False
    font = page["/Resources"]["/Font"]["/F8"].get_object()
    to_unicode = font.get("/ToUnicode")
    if to_unicode is not None:
        cmap_stream = to_unicode.get_object()
        cmap_data = cmap_stream.get_data()
        bad_digit_range = b"<056E> <0571> [<0030> <0020> <0020> <0020>]"
        correct_digit_range = b"<056E> <0571> [<0030> <0031> <0032> <0033>]"
        if bad_digit_range in cmap_data:
            if cmap_data.count(bad_digit_range) != 1:
                raise RuntimeError(f"Unsafe digit ToUnicode match count in {source}")
            cmap_stream.set_data(cmap_data.replace(bad_digit_range, correct_digit_range, 1))
            cmap_fixed = True
    writer.write(str(output))
    edits = [
        {
            "kind": "replace_credit_hours_glyph",
            "page_1_based": page_index + 1,
            "from": 4,
            "to": 2,
            "glyph_bytes": "0572 -> 0570",
            "text_matrix": matrix.decode("ascii"),
        }
    ]
    if cmap_fixed:
        edits.append(
            {
                "kind": "repair_digit_text_extraction",
                "glyph_bytes": "056E-0571",
                "unicode_values": "0030-0033",
            }
        )
    return edits


def complete_learning_hours_total(
    pdf: Path,
    *,
    page_index: int,
    rect: pymupdf.Rect,
    helpers: dict[str, Any],
    working: Path,
) -> dict[str, Any]:
    document = pymupdf.open(pdf)
    scale = helpers["insert_arabic_html"](
        document[page_index],
        rect,
        "30",
        align="center",
        font_size=14.0,
    )
    temporary = working / f"{pdf.stem}-learning-hours-total.pdf"
    document.save(temporary, garbage=4, deflate=True, clean=True)
    document.close()
    temporary.replace(pdf)
    return {
        "kind": "complete_total_learning_hours",
        "page_1_based": page_index + 1,
        "value": 30,
        "rect_top_points": [round(number, 3) for number in rect],
        "htmlbox_scale": scale,
    }


def normalize_document_properties(
    pdf: Path,
    *,
    title: str,
    code: str,
    working: Path,
    rebuild_outline: bool = False,
) -> list[dict[str, Any]]:
    """Publish an exact title and optionally replace inherited broken outlines."""
    reader = PdfReader(str(pdf), strict=False)
    writer = PdfWriter(clone_from=reader)
    subject = f"توصيف مقرر {title}"
    writer.add_metadata({"/Title": title, "/Subject": subject})
    edits = [
        {
            "kind": "set_pdf_title_metadata",
            "value": title,
            "subject": subject,
        }
    ]
    if rebuild_outline:
        writer.root_object.pop(NameObject("/Outlines"), None)
        writer.add_outline_item(f"{title} — {code}", 0)
        edits.append(
            {
                "kind": "replace_broken_outline",
                "entries": 1,
                "value": f"{title} — {code}",
            }
        )
    temporary = working / f"{pdf.stem}-document-properties.pdf"
    writer.write(str(temporary))
    temporary.replace(pdf)
    return edits


def complete_topic_hours(
    pdf: Path,
    *,
    page_index: int,
    rect: pymupdf.Rect,
    helpers: dict[str, Any],
    working: Path,
) -> dict[str, Any]:
    document = pymupdf.open(pdf)
    scale = helpers["insert_arabic_html"](
        document[page_index],
        rect,
        "ساعتان",
        align="center",
        font_size=14.0,
    )
    temporary = working / f"{pdf.stem}-topic-hours.pdf"
    document.save(temporary, garbage=4, deflate=True, clean=True)
    document.close()
    temporary.replace(pdf)
    return {
        "kind": "complete_topic_hours",
        "page_1_based": page_index + 1,
        "topic_number": 7,
        "value": "ساعتان",
        "rect_top_points": [round(number, 3) for number in rect],
        "htmlbox_scale": scale,
    }


def set_cell_text(cell: Any, value: str) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    if paragraph.runs:
        paragraph.runs[0].text = value
        for extra in paragraph.runs[1:]:
            extra.text = ""
    else:
        paragraph.add_run(value)
    for extra_paragraph in cell.paragraphs[1:]:
        for run in extra_paragraph.runs:
            run.text = ""


def complete_docx_approval(document: Document, approval: dict[str, list[str]]) -> None:
    for table in reversed(document.tables):
        if not table.rows or not any("جهة الاعتماد" in cell.text for cell in table.rows[0].cells):
            continue
        for row, key in zip(table.rows[:3], ("authority", "session", "date"), strict=True):
            value_cells = list(row.cells[1:])
            if not value_cells:
                raise RuntimeError("Approval table has no value cell")
            value = " / ".join(approval[key])
            for cell in value_cells:
                set_cell_text(cell, value)
        return
    raise RuntimeError("DOCX approval table not found")


def normalize_doctorate_docx_layout(document: Document) -> list[dict[str, Any]]:
    """Remove empty template rows and prevent known Word-to-PDF wraps."""
    outcomes = document.tables[4]
    removed = []
    for row_index in (9, 4):
        row = outcomes.rows[row_index]
        removed.append(row.cells[0].text.strip())
        outcomes._tbl.remove(row._tr)

    topics = document.tables[5]
    topics.autofit = False
    topic_widths = (Inches(0.65), Inches(4.788), Inches(1.251))
    table_properties = topics._tbl.tblPr
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    for grid_column, width in zip(
        topics._tbl.tblGrid.gridCol_lst, topic_widths, strict=True
    ):
        grid_column.w = width
    for row in topics.rows:
        grid_index = 0
        while grid_index < len(topic_widths):
            cell = row.cells[grid_index]
            span_element = cell._tc.get_or_add_tcPr().find(qn("w:gridSpan"))
            span = int(span_element.get(qn("w:val"))) if span_element is not None else 1
            cell.width = sum(topic_widths[grid_index:grid_index + span])
            grid_index += span
        number_cell = row.cells[0]
        number_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tc_pr = number_cell._tc.get_or_add_tcPr()
        if tc_pr.find(qn("w:noWrap")) is None:
            tc_pr.append(OxmlElement("w:noWrap"))
        for paragraph in number_cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(11)

    toc_entries = 0
    for paragraph in document._element.iter(qn("w:p")):
        properties = paragraph.find(qn("w:pPr"))
        style = properties.find(qn("w:pStyle")) if properties is not None else None
        if style is None or style.get(qn("w:val")) != "10":
            continue
        toc_entries += 1
        for run_properties in paragraph.iter(qn("w:rPr")):
            for tag in ("w:sz", "w:szCs"):
                element = run_properties.find(qn(tag))
                if element is None:
                    element = OxmlElement(tag)
                    run_properties.append(element)
                element.set(qn("w:val"), "26")

    return [
        {"kind": "remove_empty_learning_outcome_template_rows", "labels": removed},
        {
            "kind": "prevent_topic_number_wrapping",
            "font_size_points": 11,
            "number_column_width_inches": 0.65,
        },
        {"kind": "fit_table_of_contents", "entries": toc_entries, "font_size_points": 13},
    ]


def build_qiraat_doctorate(
    source: Path,
    output: Path,
    helpers: dict[str, Any],
    working: Path,
) -> list[dict[str, Any]]:
    document = Document(source)
    title = "دراسات في طبقات القراء وأسانيدهم"
    code = "2002703-2"
    approval = {
        "authority": ["مجلس قسم القراءات"],
        "session": ["الثالث عشر"],
        "date": ["14/9/1445هـ"],
    }
    if not helpers["replace_docx_label"](document, "اسم المقرر", title):
        raise RuntimeError("Doctorate course title field not found")
    if not helpers["replace_docx_label"](document, "رمز المقرر", "2-2002703"):
        raise RuntimeError("Doctorate course code field not found")
    layout_edits = normalize_doctorate_docx_layout(document)
    complete_docx_approval(document, approval)
    docx_output = working / "2002703-2.docx"
    document.save(docx_output)
    conversion = working / "convert-2002703-2"
    conversion.mkdir(parents=True, exist_ok=True)
    run(
        [
            str(SOFFICE),
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(conversion),
            str(docx_output),
        ],
        capture=True,
    )
    converted = conversion / "2002703-2.pdf"
    if not converted.is_file():
        raise RuntimeError(f"LibreOffice did not create {converted}")
    shutil.copy2(converted, output)
    return [
        {"kind": "replace_cover_title", "to": title},
        {"kind": "replace_course_code", "to": code},
        *layout_edits,
        {"kind": "complete_approval_table", "values": approval},
    ]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, default=OUTPUT_RELATIVE)
    parser.add_argument("--working", type=Path, default=WORKING_RELATIVE)
    args = parser.parse_args()

    root = args.root.resolve()
    output = (root / args.output).resolve()
    working = (root / args.working).resolve()
    baseline = working / "baseline"
    stamped = working / "stamped"
    for directory in (output, working, baseline, stamped):
        directory.mkdir(parents=True, exist_ok=True)

    assert_sources(root)
    helpers = import_helpers()
    records: list[dict[str, Any]] = []

    def add_record(
        *,
        filename: str,
        code: str,
        title: str,
        scope: dict[str, str],
        source: Path,
        edits: list[dict[str, Any]],
        stamp_required: bool = False,
        approval_page: int | None = None,
    ) -> None:
        path = baseline / filename
        helpers["remove_nonpublic_link_annotations"](path)
        run(["qpdf", "--check", str(path)], capture=True)
        records.append(
            {
                "filename": filename,
                "code": code,
                "title": title,
                "scope": scope,
                "source_relative_path": source.as_posix(),
                "source_sha256": sha256(root / source),
                "edits": edits,
                "stamp_required": stamp_required,
                "approval_page_1_based": approval_page,
                "baseline_sha256": sha256(path),
                "page_count_before_stamp": len(PdfReader(str(path), strict=False).pages),
            }
        )

    sira_source = root / COMBINED_ISLAMIC_STUDIES
    sira_variants = [
        (
            "2001106-2--quran.pdf",
            "بكالوريوس في القرآن وعلومه",
            {"program": "القرآن وعلومه", "degree": "بكالوريوس", "plan_type": "قديمة", "version": "39"},
            {"authority": ["مجلس قسم القراءات"], "session": ["١٥"], "date": ["١١/١١/١٤٤٥هـ"]},
        ),
        (
            "2001106-2--qiraat.pdf",
            "البكالوريوس - القراءات",
            {"program": "القراءات", "degree": "بكالوريوس", "plan_type": "قديمة", "version": "38"},
            {"authority": ["قسم القراءات"], "session": ["١"], "date": ["٧/٢/١٤٤٥هـ"]},
        ),
    ]
    for filename, program, scope, approval in sira_variants:
        path = baseline / filename
        edits = edit_sira_variant(
            sira_source,
            path,
            helpers,
            program=program,
            approval=approval,
            working=working,
        )
        add_record(
            filename=filename,
            code="2001106-2",
            title="السيرة النبوية",
            scope=scope,
            source=COMBINED_ISLAMIC_STUDIES,
            edits=edits,
            stamp_required=True,
            approval_page=9,
        )

    direct_cover_records = [
        {
            "filename": "2002235-2.pdf",
            "code": "2002235-2",
            "title": "القرآن الكريم (٧)",
            "scope": {"program": "الدراسات الإسلامية", "degree": "بكالوريوس", "plan_type": "جديدة", "version": "47"},
            "source": Path("assets/course-specifications/islamic-studies-bundle-1444/20024103-2.pdf"),
            "learning_hours_total": {
                "page_index": 3,
                "rect": pymupdf.Rect(158.2, 270.65, 241.6, 293.1),
            },
        },
        {
            "filename": "20024102-2.pdf",
            "code": "20024102-2",
            "title": "تفسير جزء عم",
            "scope": {"program": "القرآن وعلومه", "degree": "بكالوريوس", "plan_type": "جديدة", "version": "47"},
            "source": Path("assets/course-specifications/quran-old-v39/2002427-2.pdf"),
            "legacy_overlay": True,
        },
        {
            "filename": "2003347-2.pdf",
            "code": "2003347-2",
            "title": "نظام البيئة",
            "scope": {"program": "الأنظمة", "degree": "بكالوريوس", "plan_type": "جديدة", "version": "47"},
            "source": Path("assets/course-specifications/systems-source-2022/20034108-2.pdf"),
            "systems_cover": True,
        },
    ]
    for item in direct_cover_records:
        path = baseline / item["filename"]
        if item.get("legacy_overlay"):
            edits = edit_legacy_quran_cover(
                root / item["source"],
                path,
                helpers,
                title=item["title"],
                code=item["code"],
                working=working,
            )
        elif item.get("systems_cover"):
            edits = edit_systems_cover(
                root / item["source"],
                path,
                helpers,
                title=item["title"],
                code=item["code"],
            )
        else:
            edits = edit_cover_pdf(
                root / item["source"],
                path,
                helpers,
                title=item["title"],
                code=item["code"],
            )
        if item.get("learning_hours_total"):
            total = item["learning_hours_total"]
            edits.append(
                complete_learning_hours_total(
                    path,
                    page_index=total["page_index"],
                    rect=total["rect"],
                    helpers=helpers,
                    working=working,
                )
            )
        add_record(
            filename=item["filename"],
            code=item["code"],
            title=item["title"],
            scope=item["scope"],
            source=item["source"],
            edits=edits,
        )

    hours_records = [
        {
            "filename": "2002454-2--quran.pdf",
            "scope": {"program": "القرآن وعلومه", "degree": "بكالوريوس", "plan_type": "قديمة/جديدة", "version": "39/47"},
            "source": Path("assets/course-specifications/quran-old-v39/2002454-2.pdf"),
            "page_index": 1,
            "matrix": b"1 0 0 1 445.3 708.22 Tm",
            "topic_page_index": 3,
            "topic_rect": pymupdf.Rect(57.4, 274.05, 152.7, 297.7),
        },
        {
            "filename": "2002454-2--qiraat.pdf",
            "scope": {"program": "القراءات", "degree": "بكالوريوس", "plan_type": "قديمة", "version": "38"},
            "source": Path("assets/course-specifications/qiraat-source-1445/2002454-2.pdf"),
            "page_index": 2,
            "matrix": b"1 0 0 1 445.66 658.9 Tm",
            "topic_page_index": 4,
            "topic_rect": pymupdf.Rect(57.4, 639.05, 152.8, 662.55),
        },
    ]
    for item in hours_records:
        path = baseline / item["filename"]
        edits = replace_hours_glyph(
            root / item["source"],
            path,
            page_index=item["page_index"],
            matrix=item["matrix"],
        )
        edits.append(
            complete_topic_hours(
                path,
                page_index=item["topic_page_index"],
                rect=item["topic_rect"],
                helpers=helpers,
                working=working,
            )
        )
        add_record(
            filename=item["filename"],
            code="2002454-2",
            title="التفسير الموضوعي",
            scope=item["scope"],
            source=item["source"],
            edits=edits,
        )

    doctorate_source = Path(
        "ااااالملفات الخام/توصيف المقررات/"
        "توصيفات دكتوراة القراءات/"
        "دراسات في طبقات القراء وأسانيدهم.docx"
    )
    doctorate_path = baseline / "2002703-2.pdf"
    doctorate_edits = build_qiraat_doctorate(
        root / doctorate_source,
        doctorate_path,
        helpers,
        working,
    )
    add_record(
        filename="2002703-2.pdf",
        code="2002703-2",
        title="دراسات في طبقات القراء وأسانيدهم",
        scope={"program": "القراءات", "degree": "دكتوراه", "plan_type": "جديدة", "version": "1"},
        source=doctorate_source,
        edits=doctorate_edits,
        stamp_required=True,
        approval_page=len(PdfReader(str(doctorate_path), strict=False).pages),
    )

    stamp_inventory = {
        "records": [
            {
                "relative_path": record["filename"],
                "absolute_path": str(baseline / record["filename"]),
                "stamp_status": "missing_stamp",
                "department_key": "qiraat",
                "expected_stamp_asset": str(root / QIRAAT_STAMP),
                "stamp_target_page_1_based": record["approval_page_1_based"],
                "approval_location": "final_complete",
                "baseline_sha256": record["baseline_sha256"],
                "page_count": record["page_count_before_stamp"],
            }
            for record in records
            if record["stamp_required"]
        ]
    }
    inventory_path = working / "stamp-inventory.json"
    inventory_path.write_text(
        json.dumps(stamp_inventory, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    stamp_audit = output / "stamp-audit.json"
    run(
        [
            sys.executable,
            str(root / "scripts/apply_department_stamps.py"),
            "--inventory",
            str(inventory_path),
            "--output-root",
            str(stamped),
            "--audit-output",
            str(stamp_audit),
        ]
    )
    stamp_data = json.loads(stamp_audit.read_text(encoding="utf-8"))
    stamps = {item["relative_path"]: item for item in stamp_data["records"]}

    for record in records:
        filename = record["filename"]
        source_pdf = stamped / filename if record["stamp_required"] else baseline / filename
        target_pdf = output / filename
        shutil.copy2(source_pdf, target_pdf)
        helpers["add_accessible_identity"](
            target_pdf,
            record["title"],
            record["code"],
            working,
        )
        record["edits"].extend(
            normalize_document_properties(
                target_pdf,
                title=record["title"],
                code=record["code"],
                working=working,
                rebuild_outline=filename == "2002235-2.pdf",
            )
        )
        run(["qpdf", "--check", str(target_pdf)], capture=True)
        record["department_stamp"] = stamps.get(
            filename, {"status": "preserved_existing_correct_stamp"}
        )
        record["output_sha256"] = sha256(target_pdf)
        record["page_count_after_stamp"] = len(PdfReader(str(target_pdf), strict=False).pages)
        record.pop("stamp_required")
        record.pop("approval_page_1_based")

    result = {
        "schema": "course-identity-resolution-bundle-v1",
        "created_for": "approved identity and credit-hour resolutions on 2026-08-29",
        "policy": [
            "Source files remain untouched.",
            "One PDF is emitted for each program-specific identity.",
            "Ambiguous graduate Usul al-Tafsir hours/code is deliberately excluded pending a decision.",
        ],
        "counts": {
            "outputs": len(records),
            "program_specific_variants": 4,
            "credit_hours_corrected": 2,
            "stamps_added": len(stamp_inventory["records"]),
        },
        "records": records,
    }
    (output / "manifest.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (output / "hashes.sha256").write_text(
        "".join(
            f"{record['output_sha256']}  {record['filename']}\n"
            for record in records
        ),
        encoding="utf-8",
    )
    run([sys.executable, str(root / "scripts/sanitize_audit_paths.py"), str(output)])
    print(json.dumps(result["counts"], ensure_ascii=False))


if __name__ == "__main__":
    main()
