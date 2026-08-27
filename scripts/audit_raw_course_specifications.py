#!/usr/bin/env python3
"""Inventory raw course-specification sources and compare them with missing courses.

This is a read-only audit: source PDFs and Word files are never modified.  The
JSON output records hashes, extracted identities, duplicate groups, and ranked
matches against the missing-course inventory.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BASELINE = Path("/tmp/missing-specialty-inventory-20260826.json")
DEFAULT_OUTPUT = ROOT / "tmp" / "pdfs" / "raw-course-audit" / "inventory.json"

DIGIT_TRANSLATION = str.maketrans(
    "٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹",
    "01234567890123456789",
)
COURSE_CODE_PATTERN = re.compile(
    r"(?<!\d)(\d{1,2})\s*[-ـ–—]+\s*(\d{5,9})(?!\d)"
    r"|(?<!\d)(\d{5,9})\s*[-ـ–—]+\s*(\d{1,2})(?!\d)"
)
TITLE_CODE_PATTERN = re.compile(
    r"اسم\s+(?:ال|ا?مل)?مقرر\s*[:：]?\s*(?P<title>.*?)\s*"
    r"رمز\s+(?:ال|ا?مل)?مقرر\s*[:：]?\s*(?P<code>.{0,100})",
    re.DOTALL,
)
PROGRAM_PATTERN = re.compile(
    r"(?:البرنامج|اسم\s+البرنامج)\s*[:：]?\s*(.*?)\s*"
    r"(?:القسم\s+العلمي|الكلية|المؤسسة|نسخة\s+التوصيف|$)",
    re.DOTALL,
)
GENERIC_FILENAME_WORDS = {
    "توصيف",
    "توصيفات",
    "مقرر",
    "المقرر",
    "مادة",
    "نموذج",
    "اصدار",
    "برنامج",
    "كلية",
    "الشريعة",
    "الانظمة",
    "pdf",
    "docx",
    "نسخة",
}


def normalize_digits(value: str) -> str:
    return (value or "").translate(DIGIT_TRANSLATION)


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_arabic(value: str, *, drop_article: bool = True) -> str:
    text = unicodedata.normalize("NFKC", normalize_digits(value or "")).lower()
    text = re.sub(r"[\u061c\u200e\u200f\u202a-\u202e\u2066-\u2069]", "", text)
    text = text.replace("ـ", "")
    text = "".join(char for char in unicodedata.normalize("NFD", text) if unicodedata.category(char) != "Mn")
    text = unicodedata.normalize("NFC", text)
    text = re.sub(r"[أإآٱ]", "ا", text)
    text = re.sub(r"[ؤئ]", "ء", text)
    text = re.sub(r"[ىيی]", "ي", text)
    text = re.sub(r"[ةۀہ]", "ه", text)
    text = re.sub(r"[كک]", "ك", text)
    text = re.sub(r"([^\W\d_])(\d)", r"\1 \2", text, flags=re.UNICODE)
    text = re.sub(r"(\d)([^\W\d_])", r"\1 \2", text, flags=re.UNICODE)
    tokens = re.findall(r"[^\W_]+", text, flags=re.UNICODE)
    if drop_article:
        tokens = [token[2:] if token.startswith("ال") and len(token) > 4 else token for token in tokens]
    return " ".join(token for token in tokens if token)


def clean_title(value: str) -> str:
    value = normalize_space(value)
    value = re.sub(r"^[\s:：\-–—.،؛]+|[\s:：\-–—.،؛]+$", "", value)
    return normalize_space(value)


def parse_course_code(value: str) -> str | None:
    normalized = normalize_digits(value or "")
    for match in COURSE_CODE_PATTERN.finditer(normalized):
        short_left, long_right, long_left, short_right = match.groups()
        if long_left and short_right:
            return f"{long_left}-{short_right}"
        if long_right and short_left:
            return f"{long_right}-{short_left}"
    return None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_pdf_text(path: Path) -> tuple[str, int]:
    reader = PdfReader(str(path), strict=False)
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as error:  # retain the usable pages and record the warning in text
            parts.append(f"\n[PAGE_EXTRACTION_ERROR: {error}]\n")
    return "\n".join(parts), len(reader.pages)


def filename_title(path: Path) -> str:
    value = normalize_arabic(path.stem, drop_article=False)
    tokens = [
        token
        for token in value.split()
        if token not in GENERIC_FILENAME_WORDS
        and not re.fullmatch(r"(?:14|20)\d{2}", token)
        and not re.fullmatch(r"\d{6,9}", token)
    ]
    while tokens and tokens[-1] in {"1", "2", "3"} and "(" in path.stem:
        break
    return normalize_space(" ".join(tokens))


def parse_identities(text: str, path: Path) -> list[dict[str, Any]]:
    normalized_labels = text.replace("املقرر", "المقرر").replace("ا لمقرر", "المقرر")
    matches = list(TITLE_CODE_PATTERN.finditer(normalized_labels))
    identities: list[dict[str, Any]] = []

    for index, match in enumerate(matches):
        title = clean_title(match.group("title").splitlines()[-1])
        code_field = match.group("code").splitlines()[0]
        code = parse_course_code(code_field)
        if not code:
            code = parse_course_code(match.group(0))
        if not title or len(title) > 180:
            continue

        segment_end = matches[index + 1].start() if index + 1 < len(matches) else min(len(normalized_labels), match.end() + 1200)
        segment = normalized_labels[match.end():segment_end]
        program_match = PROGRAM_PATTERN.search(segment)
        program = clean_title(program_match.group(1).splitlines()[0]) if program_match else None
        identities.append(
            {
                "title": title,
                "code": code,
                "raw_code_field": normalize_space(code_field),
                "program": program,
                "source": "document_fields",
            }
        )

    if identities:
        return identities

    fallback_code = parse_course_code(path.stem) or parse_course_code(normalized_labels[:4000])
    fallback_title = filename_title(path)
    if fallback_title:
        identities.append(
            {
                "title": fallback_title,
                "code": fallback_code,
                "raw_code_field": None,
                "program": None,
                "source": "filename_fallback",
            }
        )
    return identities


def title_score(left: str, right: str) -> float:
    a = normalize_arabic(left)
    b = normalize_arabic(right)
    if not a or not b:
        return 0.0
    a_numbers = re.findall(r"\d+", a)
    b_numbers = re.findall(r"\d+", b)
    if a_numbers and b_numbers and a_numbers != b_numbers:
        return 0.0
    if a == b:
        return 1.0
    a_tokens = set(a.split())
    b_tokens = set(b.split())
    overlap = len(a_tokens & b_tokens)
    token_f1 = (2 * overlap / (len(a_tokens) + len(b_tokens))) if overlap else 0.0
    sequence = SequenceMatcher(None, a, b).ratio()
    containment = min(len(a), len(b)) / max(len(a), len(b)) if a in b or b in a else 0.0
    return max(token_f1, sequence, containment)


def missing_items(baseline: dict[str, Any]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for raw in baseline.get("missing_no_source_or_candidate", []):
        occurrences = raw.get("occurrences") or []
        title = raw.get("title") or raw.get("name") or (occurrences[0].get("title") if occurrences else "")
        result.append(
            {
                "identity_key": raw.get("identity_key") or (occurrences[0].get("identity_key") if occurrences else None),
                "code": raw.get("code") or (occurrences[0].get("code") if occurrences else None),
                "title": title,
                "occurrences": occurrences,
            }
        )
    return result


def classify_match(target: dict[str, Any], candidate: dict[str, Any]) -> tuple[str | None, float]:
    candidate_code = candidate.get("code")
    target_code = target.get("code")
    score = title_score(target.get("title", ""), candidate.get("title", ""))
    code_equal = bool(candidate_code and target_code and candidate_code == target_code)

    if code_equal and score >= 0.94:
        return "exact_identity", 100.0 + score
    if code_equal and score >= 0.70:
        return "code_match_near_title", 92.0 + score
    if code_equal:
        return "code_only_conflict", 86.0 + score
    if score >= 0.97 and candidate_code:
        return "title_match_code_conflict", 82.0 + score
    if score >= 0.97:
        return "title_match_no_code", 78.0 + score
    if score >= 0.86:
        return "near_title", 70.0 + score
    return None, score


def inventory_file(path: Path, raw_root: Path) -> dict[str, Any]:
    record: dict[str, Any] = {
        "path": path.relative_to(ROOT).as_posix(),
        "raw_relative_path": path.relative_to(raw_root).as_posix(),
        "extension": path.suffix.lower(),
        "size": path.stat().st_size,
        "sha256": sha256(path),
        "pages": None,
        "text_length": 0,
        "text_error": None,
        "identities": [],
    }
    try:
        if path.suffix.lower() == ".pdf":
            text, pages = extract_pdf_text(path)
            record["pages"] = pages
        else:
            text = extract_docx_text(path)
        record["text_length"] = len(text)
        record["identities"] = parse_identities(text, path)
    except Exception as error:
        record["text_error"] = f"{type(error).__name__}: {error}"
        record["identities"] = parse_identities("", path)
    return record


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-root", type=Path)
    parser.add_argument("--baseline", type=Path, default=DEFAULT_BASELINE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--include", choices=("pdf", "docx", "both"), default="both")
    args = parser.parse_args()

    raw_root = args.raw_root
    if raw_root is None:
        candidates = sorted(path for path in ROOT.iterdir() if path.is_dir() and "الملفات الخام" in path.name)
        if len(candidates) != 1:
            raise RuntimeError(f"Expected exactly one raw-files directory, found {len(candidates)}")
        raw_root = candidates[0]
    raw_root = raw_root.resolve()

    suffixes = {".pdf", ".docx"} if args.include == "both" else {f".{args.include}"}
    paths = sorted(
        (path for path in raw_root.rglob("*") if path.is_file() and path.suffix.lower() in suffixes),
        key=lambda path: path.as_posix(),
    )
    files: list[dict[str, Any]] = []
    for index, path in enumerate(paths, start=1):
        files.append(inventory_file(path, raw_root))
        if index % 50 == 0 or index == len(paths):
            print(f"inventoried {index}/{len(paths)}", flush=True)

    hash_groups: defaultdict[str, list[str]] = defaultdict(list)
    for record in files:
        hash_groups[record["sha256"]].append(record["path"])

    candidate_records: list[dict[str, Any]] = []
    for file_record in files:
        for identity_index, identity in enumerate(file_record["identities"]):
            candidate_records.append(
                {
                    **identity,
                    "path": file_record["path"],
                    "extension": file_record["extension"],
                    "sha256": file_record["sha256"],
                    "duplicate_paths": hash_groups[file_record["sha256"]],
                    "identity_index": identity_index,
                }
            )

    baseline = json.loads(args.baseline.read_text(encoding="utf-8"))
    targets = missing_items(baseline)
    matched_targets: list[dict[str, Any]] = []
    for target in targets:
        ranked: list[dict[str, Any]] = []
        seen: set[tuple[str, str | None, str]] = set()
        for candidate in candidate_records:
            category, rank = classify_match(target, candidate)
            if not category:
                continue
            key = (candidate["sha256"], candidate.get("code"), normalize_arabic(candidate.get("title", "")))
            if key in seen:
                continue
            seen.add(key)
            ranked.append(
                {
                    "category": category,
                    "rank": round(rank, 6),
                    "title_score": round(title_score(target["title"], candidate.get("title", "")), 6),
                    **candidate,
                }
            )
        ranked.sort(
            key=lambda item: (
                -item["rank"],
                0 if item["extension"] == ".pdf" else 1,
                item["path"],
            )
        )
        matched_targets.append({**target, "candidates": ranked[:20]})

    category_counts = Counter(
        item["candidates"][0]["category"] if item["candidates"] else "no_candidate"
        for item in matched_targets
    )
    result = {
        "schema": 1,
        "source_root": raw_root.as_posix(),
        "baseline": args.baseline.as_posix(),
        "summary": {
            "files": len(files),
            "pdf_files": sum(record["extension"] == ".pdf" for record in files),
            "docx_files": sum(record["extension"] == ".docx" for record in files),
            "unique_hashes": len(hash_groups),
            "duplicate_files": len(files) - len(hash_groups),
            "text_errors": sum(bool(record["text_error"]) for record in files),
            "parsed_identities": len(candidate_records),
            "missing_targets": len(targets),
            "best_match_categories": dict(sorted(category_counts.items())),
        },
        "duplicate_groups": [
            {"sha256": digest, "paths": paths}
            for digest, paths in sorted(hash_groups.items())
            if len(paths) > 1
        ],
        "files": files,
        "matches": matched_targets,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result["summary"], ensure_ascii=False, indent=2))
    print(f"wrote {args.output}")


if __name__ == "__main__":
    main()
