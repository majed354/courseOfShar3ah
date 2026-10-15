"""Reproduce the authorized Master CLO revision from preserved source PDFs and decisions.

Run prepare_quranic_master_source_pages_20261015.py first; then build documents
with the bundled document Python. PDF export uses the bundled renderer.
"""
import json,copy,re
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

R=Path(__file__).resolve().parents[1];B=R/'assets/course-specifications/quranic-master-single-domain-20261015';Q=R/'.netlify/qa/master_single_domain_20261015';cfg=json.loads((B/'revision-decisions.json').read_text());ims=json.loads((Q/'course-image-manifest.json').read_text());PURPLE='4E3F8C';TEAL='52B4C0'
def rtl(p):
 pr=p._p.get_or_add_pPr();b=OxmlElement('w:bidi');pr.append(b);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.space_after=Pt(5)
def run(p,text,bold=False,size=14,color=None):
 r=p.add_run(str(text));r.font.name='Sakkal Majalla';r.font.size=Pt(size);r.bold=bold
 f=r._r.get_or_add_rPr().get_or_add_rFonts()
 for k in ['ascii','hAnsi','cs']:f.set(qn('w:'+k),'Sakkal Majalla')
 sz=OxmlElement('w:szCs');sz.set(qn('w:val'),str(size*2));r._r.rPr.append(sz)
 if color:r.font.color.rgb=RGBColor.from_string(color)
 return r
def para(doc,text,bold=False,size=14,color=None):
 p=doc.add_paragraph();rtl(p);run(p,text,bold,size,color);p.paragraph_format.keep_with_next=bold;return p
def put(cell,text,center=False,size=14):
 cell.text='';p=cell.paragraphs[0];rtl(p);p.paragraph_format.space_after=Pt(2)
 if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 run(p,text,size=size)
def table(doc,headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'FFFFFF');borders.append(x)
 t._tbl.tblPr.append(borders)
 for i,(c,h,w) in enumerate(zip(t.rows[0].cells,headers,widths)):
  c.width=Cm(w);put(c,h,True);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),PURPLE);c._tc.get_or_add_tcPr().append(shade)
  for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)

 for col,w in zip(t.columns,widths):col.width=Cm(w)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for j,vals in enumerate(rows):
  row=t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c,value,w in zip(row.cells,vals,widths):
   c.width=Cm(w);put(c,value);s=OxmlElement('w:shd');s.set(qn('w:fill'),'F1F1F1' if j%2==0 else 'DDDDDD');c._tc.get_or_add_tcPr().append(s)
 return t
def newdoc():
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.5);s.bottom_margin=Cm(1.5);s.left_margin=Cm(1.6);s.right_margin=Cm(1.6);s.footer_distance=Cm(.5)
 p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
 return d
def fullpage(doc,path):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
 inline=p.add_run().add_picture(path,width=Cm(21),height=Cm(29.7))._inline;anchor=OxmlElement('wp:anchor')
 for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items():anchor.set(k,v)
 simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
 for tag in ['positionH','positionV']:
  e=OxmlElement('wp:'+tag);e.set('relativeFrom','page');o=OxmlElement('wp:posOffset');o.text='0';e.append(o);anchor.append(e)
 for name in ['extent','effectExtent']:
  e=inline.find(qn('wp:'+name))
  if e is not None:anchor.append(copy.deepcopy(e))
 anchor.append(OxmlElement('wp:wrapNone'))
 for name in ['docPr','cNvGraphicFramePr']:
  e=inline.find(qn('wp:'+name))
  if e is not None:anchor.append(copy.deepcopy(e))
 anchor.append(copy.deepcopy(inline.find(qn('a:graphic'))));inline.getparent().replace(inline,anchor)
def rubric(doc):
 r=cfg['q2_project_rubric'];para(doc,'سلم تقويم المخرج القيمي 3.2 في المشروع البحثي',True,18,PURPLE)
 para(doc,'ناتج البرنامج ق2 — المسؤولية والاستقلالية والتعاون والقيادة البحثية')
 table(doc,['المعيار','الوزن الداخلي','الدليل المطلوب'],[[x['name'],str(x['weight'])+'%',x['evidence']] for x in r['criteria']],[6.5,2,9.3])
 table(doc,['الدرجة من4','وصف الأداء'],list(r['levels'].items()),[2,15.8]);para(doc,r['score_formula'],True)
 for text in r['rules']:para(doc,text)
 para(doc,'اسم الطالب: ................................  رمز المقرر: 2002888-4  العام: ..........')
 para(doc,'درجات المعايير الأربعة: .......... / .......... / .......... / ..........   النتيجة: .......... /100')
 para(doc,'أسماء الأدلة وتواريخها: ........................................................................')
 para(doc,'اسم المقوّم وتاريخ التقويم: ..................................................................')
for c in cfg['courses']:
 d=newdoc();meta=ims[c['code']];d.add_picture(meta['header'],width=Cm(17.8));para(d,'توصيف المقرر الدراسي',True,24,PURPLE);para(d,c['name'],True,22,TEAL)
 table(d,['الحقل','البيان'],[['رمز المقرر',c['code']],['البرنامج','ماجستير الدراسات القرآنية المعاصرة'],['الساعات المعتمدة',c['code'].split('-')[1]],['المستوى',c['level']],['المراجعة','15 أكتوبر 2026 — الربط الأحادي داخل المجال']],[4.2,13.6])
 para(d,'تستبدل هذه المراجعة جدول نواتج تعلم المقرر وربطه السابق. يرتبط كل مخرج بناتج برنامج واحد ضمن مجاله: 1 معرفي، 2 مهاري، 3 قيمي. تبقى موضوعات المقرر ومصادره وساعاته وأوزان أنشطة التقويم وفق المصدر المحفوظ.')
 para(d,'اعتماد التوصيف المحدّث',True,17,PURPLE);table(d,['جهة الاعتماد','رقم الجلسة','تاريخ الجلسة'],[['مجلس قسم القراءات','','']],[8.8,4.5,4.5])
 for image in meta['before']:d.add_page_break();fullpage(d,image)
 d.add_page_break();d.add_picture(meta['header'],width=Cm(17.8));para(d,'نواتج تعلم المقرر واستراتيجيات تدريسها وطرق تقويمها',True,18,PURPLE)
 table(d,['الرمز','ناتج تعلم المقرر','PLO','التدريس','التقويم'],[[x['code'],x['text'],x['plo_code'],x['teaching'],x['assessment']] for x in c['clos']],[1.4,7.1,1.3,3.3,4.7])
 para(d,'أوزان أنشطة تقويم المقرر',True,17,PURPLE)
 table(d,['نشاط التقويم','النسبة من الدرجة'],[[a['label'],str(a['weight'])+'%'] for a in c['assessment_plan_preserved']],[13.8,4])
 para(d,'تسجل درجة مستقلة لكل مخرج وفق مؤشرات أدائه. سلالم التقدير تفصل أداء المهارة عن القيم، وتستخدم ضمن أنشطة التقويم القائمة دون زيادة مجموع أوزان المقرر عن100%.')
 if c['code']=='2002888-4':d.add_page_break();rubric(d)
 for image in meta['after']:d.add_page_break();fullpage(d,image)
 folder=B/'courses'/c['code'];folder.mkdir(parents=True,exist_ok=True);d.save(folder/'course-specification.docx')
 print('Built',c['code'])
# The official programme template is retained; edit the matrix and crosswalk.
source=R/'assets/course-specifications/quranic-master-official-template-20261015/program-specification.docx';d=Document(source)
for row,c in zip(d.tables[9].rows[3:],cfg['courses']):
 for i,o in enumerate(cfg['outcomes'],1):put(row.cells[i],c['matrix'].get(o['code'],''),True,12)
t=d.tables[31];prototype=copy.deepcopy(t.rows[1]._tr)
for row in list(t.rows)[1:]:t._tbl.remove(row._tr)
for c in cfg['courses']:
 for x in c['clos']:
  tr=copy.deepcopy(prototype);t._tbl.append(tr);row=t.rows[-1]
  for cell,value in zip(row.cells,[c['code'],c['name'],x['code'],x['plo_code']]):put(cell,value,True,12)
for rel in d.part.rels.values():
 if 'hyperlink' in rel.reltype:
  for c in cfg['courses']:
   if rel.target_ref.endswith('/'+c['code']+'.pdf'):
    rel._target='https://courseofshar3ah.netlify.app/assets/course-specifications/quranic-master-single-domain-20261015/courses/'+c['code']+'/course-specification.pdf'
for p in d.paragraphs:
 if 'تبقى نصوص نواتج المقررات وخطط تقويمها' in p.text:
  p.clear();rtl(p);run(p,'يحدد هذا الملحق الربط المحدث في سياق ماجستير الدراسات القرآنية المعاصرة، ويحل محل الربط السابق. عُدلت صياغات النواتج وتصنيفاتها وأدوات تقويمها حيث لزم، وأضيف المخرج القيمي3.2 في المشروع البحثي. بقيت الخطة والمقررات والساعات وموضوعاتها وأوزان أنشطة تقويمها ثابتة. الرمز1 تأسيس، و2 ممارسة، و3 تمكين؛ ولا يعني الربط وحده توافر نتيجة قياس.')
 if 'يقاس ق2' in p.text or 'يُقاس ق2' in p.text:
  p.clear();rtl(p);run(p,'يقاس ق2 في منهج البحث والتحقيق من المخرجين القيميين 3.1 و3.3، وفي المشروع البحثي من المخرج القيمي المستقل 3.2. تستخدم درجة مستقلة وسجل أدلة وسلم تقويم للمسؤولية والاستقلالية والمبادرة والتعاون والقيادة. لا تُنقل درجة المخرج المهاري2.2 أو درجة النزاهة3.1 أو الدرجة الإجمالية للمشروع إلى ق2. لا يعاد ترميز نتائج سابقة دون قياس موثق للمخرج الجديد.')
 if 'البنود غير المرمزة' in p.text:
  p.clear();rtl(p);run(p,'حُسمت البنود الثلاثة غير المرمزة في هذه المراجعة، وأعيدت صياغة النواتج وتصنيفها وفق مجالها. يعرض الملحق94 مخرج مقرر؛ لكل منها ناتج برنامج واحد في المجال نفسه. روابط التصنيف3 وحدها تمثل التمكين، ويظل توفر الربط مستقلًا عن توفر نتيجة القياس.')
d.add_page_break();para(d,'ملحق3 سلم تقويم المسؤولية والاستقلالية في المشروع البحثي',True,18,PURPLE);rubric(d)
# No historic council approval is carried into this authored revision.
put(d.tables[30].rows[1].cells[1],'');put(d.tables[30].rows[2].cells[1],'')
# Replace stale cached page numbers and keep the index off the cover.
for content in d._element.iter(qn('w:sdtContent')):
 if not any('جدول المحتويات' in (t.text or '') for t in content.iter(qn('w:t'))):continue
 for child in list(content):content.remove(child)
 for i,(label,page) in enumerate([('جدول المحتويات',''),('أ. التعريف بالبرنامج ومعلومات عامة عنه',3),('ب. رسالة البرنامج وأهدافه ونواتج التعلم',4),('ج. المنهج الدراسي',5),('د. الرسالة العلمية ومتطلباتها',10),('هـ. القبول والدعم الطلابي',10),('و. هيئة التدريس والموظفون',11),('ز. مصادر التعلم والمرافق والتجهيزات',12),('ح. ضمان جودة البرنامج',13),('ط. اعتماد التوصيف',17),('ملحق1: ربط نواتج المقررات بنواتج البرنامج',18),('ملحق2: أدلة التمكين ومؤشرات التقويم',21),('ملحق3: سلم تقويم ق2 في المشروع البحثي',22)]):
  node=OxmlElement('w:p');content.append(node);p=Paragraph(node,d._body);rtl(p);p.paragraph_format.space_after=Pt(9)
  if i==0:p.paragraph_format.page_break_before=True
  run(p,label+('  ........................  '+str(page) if page else ''),bold=i==0,size=14)
 break

# Preserve every untouched OOXML package member.
settings=d.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
replacements={'word/document.xml':etree.tostring(d._element,xml_declaration=True,encoding='UTF-8',standalone=True),'word/settings.xml':etree.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True),'word/_rels/document.xml.rels':d.part.rels.xml}
with ZipFile(source) as original,ZipFile(B/'program-specification.docx','w') as final:
 for info in original.infolist():final.writestr(info,replacements.get(info.filename,original.read(info.filename)))
d=newdoc();rubric(d);d.save(B/'q2-project-rubric.docx')
print('Built program and rubric')
