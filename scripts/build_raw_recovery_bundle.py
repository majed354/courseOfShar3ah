#!/usr/bin/env python3
"""Build publication PDFs for courses missing from the site and found in raw files.

The audited source manifest is deliberately strict: existing site specifications
are excluded, exact byte duplicates are collapsed, and identity conflicts are
not published.  Raw sources are never edited in place.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

import pdfplumber
import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, NameObject


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_RELATIVE = Path("assets/course-specifications/raw-recovery-20260827")
WORKING_RELATIVE = Path("tmp/pdfs/raw-recovery-20260827")
TEMP_MANIFEST = Path("/tmp/raw-recovery-manifest-final.json")
ARIAL = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
SOFFICE = Path(
    "/Users/majd/.cache/codex-runtimes/codex-primary-runtime/"
    "dependencies/bin/override/soffice"
)


def run(command: list[str], *, capture: bool = False) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        command,
        check=True,
        text=True,
        stdout=subprocess.PIPE if capture else None,
        stderr=subprocess.PIPE if capture else None,
    )


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def remove_nonpublic_link_annotations(path: Path) -> list[dict[str, Any]]:
    """Remove inherited blank/local-machine links without changing page content."""

    reader = PdfReader(str(path), strict=False)
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    removed: list[dict[str, Any]] = []

    for page_number, page in enumerate(writer.pages, start=1):
        annotations = page.get("/Annots")
        if not annotations:
            continue
        kept = ArrayObject()
        for annotation_reference in annotations:
            annotation = annotation_reference.get_object()
            action = annotation.get("/A") if annotation.get("/Subtype") == "/Link" else None
            uri = str(action.get("/URI", "")) if action and action.get("/S") == "/URI" else ""
            normalized_uri = uri.strip().lower()
            if normalized_uri == "about:blank" or normalized_uri.startswith("file:///c:/"):
                removed.append({"page_1_based": page_number, "uri": uri})
                continue
            kept.append(annotation_reference)
        if len(kept) == len(annotations):
            continue
        if kept:
            page[NameObject("/Annots")] = kept
        else:
            del page[NameObject("/Annots")]

    if not removed:
        return []

    temporary = path.with_suffix(".links-cleaned.pdf")
    with temporary.open("wb") as stream:
        writer.write(stream)
    temporary.replace(path)
    return [
        {
            "kind": "remove_nonpublic_link_annotations",
            "count": len(removed),
            "pages_1_based": sorted({item["page_1_based"] for item in removed}),
            "targets": sorted({item["uri"] for item in removed}),
        }
    ]


def required_edit(record: dict[str, Any], kind: str) -> dict[str, Any] | None:
    return next(
        (item for item in record["required_edits"] if item["kind"] == kind),
        None,
    )


def normalize_page_range(value: str, page_count: int) -> str:
    if not value:
        return f"1-{page_count}"
    return value.replace(" ", "")


def extract_pdf_pages(source: Path, pages: str, output: Path) -> None:
    source_count = len(PdfReader(str(source), strict=False).pages)
    normalized = normalize_page_range(pages, source_count)
    if normalized == f"1-{source_count}":
        shutil.copy2(source, output)
        return
    run(
        [
            "qpdf",
            "--empty",
            "--pages",
            str(source),
            normalized,
            "--",
            str(output),
        ]
    )


def most_common_background(page: pymupdf.Page, rect: pymupdf.Rect) -> tuple[float, float, float]:
    pixmap = page.get_pixmap(clip=rect, colorspace=pymupdf.csRGB, alpha=False)
    pixels = zip(pixmap.samples[0::3], pixmap.samples[1::3], pixmap.samples[2::3])
    quantized = Counter((r // 8, g // 8, b // 8) for r, g, b in pixels)
    red, green, blue = quantized.most_common(1)[0][0]
    return ((red * 8 + 4) / 255, (green * 8 + 4) / 255, (blue * 8 + 4) / 255)


def insert_arabic_html(
    page: pymupdf.Page,
    rect: pymupdf.Rect,
    text: str,
    *,
    align: str = "right",
    font_size: float = 14.0,
    color: str = "#000000",
) -> float:
    css = f"""
        @font-face {{ font-family: SiteArabic; src: url('{ARIAL.name}'); }}
        html, body {{ margin: 0; padding: 0; }}
        div {{
            margin: 0; padding: 0;
            font-family: SiteArabic;
            font-size: {font_size}pt;
            color: {color};
            direction: rtl;
            text-align: {align};
            line-height: 1.05;
        }}
    """
    spare_height, scale = page.insert_htmlbox(
        rect,
        f'<div dir="rtl">{html.escape(text)}</div>',
        css=css,
        archive=pymupdf.Archive(str(ARIAL.parent)),
        scale_low=0.55,
        overlay=True,
    )
    if spare_height < 0:
        raise RuntimeError(f"Arabic replacement did not fit: {text}")
    return float(scale)


def cover_value_rect(page: pymupdf.Page, kind: str) -> pymupdf.Rect:
    groups: dict[int, list[tuple[Any, ...]]] = {}
    for word in page.get_text("words"):
        groups.setdefault(int(word[5]), []).append(word)

    for words in groups.values():
        tokens = [str(word[4]) for word in words]
        has_course_label = any("قرر" in token for token in tokens)
        has_kind = any(("اسم" if kind == "title" else "رمز") in token for token in tokens)
        if not has_course_label or not has_kind:
            continue
        label_words = [
            word
            for word in words
            if "قرر" in str(word[4])
            or ("اسم" if kind == "title" else "رمز") in str(word[4])
        ]
        cutoff = min(float(word[0]) for word in label_words) - 3.0
        value_words = [word for word in words if float(word[2]) <= cutoff + 1.0]
        if not value_words:
            continue
        top = min(float(word[1]) for word in words) - 1.5
        bottom = max(float(word[3]) for word in words) + 1.5
        value_left = min(float(word[0]) for word in value_words) - 4.0
        left = max(55.0, min(value_left, 220.0 if kind == "title" else value_left))
        return pymupdf.Rect(left, top, cutoff, bottom)
    raise RuntimeError(f"Could not locate cover {kind} field")


def replace_cover_value(page: pymupdf.Page, kind: str, value: str) -> dict[str, Any]:
    rect = cover_value_rect(page, kind)
    background = most_common_background(page, rect)
    page.add_redact_annot(rect, fill=background, cross_out=False)
    page.apply_redactions(
        images=pymupdf.PDF_REDACT_IMAGE_NONE,
        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
        text=pymupdf.PDF_REDACT_TEXT_REMOVE,
    )
    scale = insert_arabic_html(
        page,
        rect,
        value,
        align="right",
        font_size=14.0 if kind == "title" else 13.0,
    )
    return {
        "kind": f"replace_cover_{kind}",
        "rect_top_points": [round(number, 3) for number in rect],
        "htmlbox_scale": scale,
    }


def replace_credit_hours(
    document: pymupdf.Document, page_number: int, old_hours: int, new_hours: int
) -> dict[str, Any]:
    old_text = "ساعتان" if old_hours == 2 else "أربع ساعات"
    new_text = "ساعتان" if new_hours == 2 else "(أربع ساعات)"
    page = document[page_number - 1]
    matches = page.search_for(old_text)
    if not matches:
        raise RuntimeError(f"Credit-hour phrase not found on page {page_number}: {old_text}")
    union = pymupdf.Rect(matches[0])
    for match in matches[1:]:
        union |= match
    left_extension = 75.0 if len(new_text) > len(old_text) else 6.0
    rect = pymupdf.Rect(
        max(55.0, union.x0 - left_extension),
        union.y0 - 2.0,
        min(page.rect.width - 40.0, union.x1 + 4.0),
        union.y1 + 2.0,
    )
    background = most_common_background(page, rect)
    page.add_redact_annot(rect, fill=background, cross_out=False)
    page.apply_redactions(
        images=pymupdf.PDF_REDACT_IMAGE_NONE,
        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
        text=pymupdf.PDF_REDACT_TEXT_REMOVE,
    )
    scale = insert_arabic_html(
        page,
        rect,
        new_text,
        font_size=13.0,
        color="#ffffff" if new_hours == 4 else "#000000",
    )
    return {
        "kind": "replace_credit_hours",
        "page_1_based": page_number,
        "from": old_hours,
        "to": new_hours,
        "rect_top_points": [round(number, 3) for number in rect],
        "htmlbox_scale": scale,
    }


def approval_value_rects(path: Path, page_number: int) -> list[pymupdf.Rect]:
    with pdfplumber.open(path) as pdf:
        tables = pdf.pages[page_number - 1].find_tables()
        candidates = [
            table
            for table in tables
            if len(table.extract()) == 3 and table.bbox[2] - table.bbox[0] > 400
        ]
        if not candidates:
            raise RuntimeError(f"Approval table not found: {path} page {page_number}")
        table = max(candidates, key=lambda item: item.bbox[1])
        row_zero = [
            cell
            for cell in table.cells
            if abs(cell[1] - table.bbox[1]) < 2.0 and cell[2] - cell[0] > 50
        ]
        row_zero.sort(key=lambda item: item[0])
        if len(row_zero) < 2:
            raise RuntimeError(f"Approval table columns not found: {path}")
        value_column = row_zero[0]
        row_edges = sorted(
            {
                round(value, 3)
                for cell in table.cells
                for value in (cell[1], cell[3])
                if table.bbox[1] - 2 <= value <= table.bbox[3] + 2
            }
        )
        clustered: list[float] = []
        for value in row_edges:
            if not clustered or value - clustered[-1] > 3.0:
                clustered.append(value)
        if len(clustered) != 4:
            clustered = [
                table.bbox[1] + (table.bbox[3] - table.bbox[1]) * index / 3
                for index in range(4)
            ]
        return [
            pymupdf.Rect(
                value_column[0] + 4.0,
                clustered[index] + 2.0,
                value_column[2] - 4.0,
                clustered[index + 1] - 2.0,
            )
            for index in range(3)
        ]


def complete_pdf_approval(
    path: Path, page_number: int, values: dict[str, list[str]]
) -> list[dict[str, Any]]:
    rects = approval_value_rects(path, page_number)
    document = pymupdf.open(path)
    page = document[page_number - 1]
    inserted: list[dict[str, Any]] = []
    rows = ["authority", "session", "date"]
    for row, rect in zip(rows, rects, strict=True):
        text = " / ".join(values[row])
        background = most_common_background(page, rect)
        page.add_redact_annot(rect, fill=background, cross_out=False)
        page.apply_redactions(
            images=pymupdf.PDF_REDACT_IMAGE_NONE,
            graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
            text=pymupdf.PDF_REDACT_TEXT_REMOVE,
        )
        scale = insert_arabic_html(page, rect, text, align="center", font_size=10.5)
        inserted.append(
            {
                "row": row,
                "value": text,
                "rect_top_points": [round(number, 3) for number in rect],
                "htmlbox_scale": scale,
            }
        )
    temporary = path.with_name(path.stem + "-approval.pdf")
    document.save(temporary, garbage=4, deflate=True, clean=True)
    document.close()
    temporary.replace(path)
    return inserted


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run_item in paragraph.runs[1:]:
            run_item.text = ""
    else:
        paragraph.add_run(text)
    for extra in cell.paragraphs[1:]:
        for run_item in extra.runs:
            run_item.text = ""


def replace_docx_label(document: Document, label: str, value: str) -> bool:
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if label in cell.text:
                    set_cell_text(cell, f"{label}:  {value}")
                    return True
    return False


def rtl_docx_course_code(code: str) -> str:
    """Store the code in the logical order that Word renders as base-suffix in RTL."""
    base, suffix = code.rsplit("-", 1)
    return f"{suffix}-{base}"


def complete_docx_approval(document: Document, values: dict[str, list[str]]) -> None:
    for table in reversed(document.tables):
        if not table.rows or not any("جهة الاعتماد" in cell.text for cell in table.rows[0].cells):
            continue
        row_values = [values["authority"], values["session"], values["date"]]
        for row, provided in zip(table.rows[:3], row_values, strict=True):
            value_cells = list(row.cells[1:])
            if not value_cells:
                raise RuntimeError("Approval table has no value cell")
            if len(value_cells) == 1 and len(provided) > 1:
                set_cell_text(value_cells[0], " / ".join(provided))
                continue
            expanded = provided if len(provided) > 1 else provided * len(value_cells)
            for cell, value in zip(value_cells, expanded, strict=False):
                set_cell_text(cell, value)
        return
    raise RuntimeError("DOCX approval table not found")


def convert_docx(record: dict[str, Any], output: Path, directory: Path) -> list[dict[str, Any]]:
    source = Path(record["source"]["absolute_path"])
    document = Document(source)
    edits: list[dict[str, Any]] = []
    title_edit = required_edit(record, "replace_cover_title")
    if title_edit:
        if not replace_docx_label(document, "اسم المقرر", title_edit["to_value"]):
            raise RuntimeError(f"DOCX title field not found: {source}")
        edits.append({"kind": "replace_cover_title", "to": title_edit["to_value"]})
    code_edit = required_edit(record, "replace_course_code")
    if code_edit:
        if not replace_docx_label(
            document,
            "رمز المقرر",
            rtl_docx_course_code(code_edit["to_value"]),
        ):
            raise RuntimeError(f"DOCX code field not found: {source}")
        edits.append({"kind": "replace_course_code", "to": code_edit["to_value"]})
    approval_edit = required_edit(record, "complete_approval_table")
    if approval_edit:
        complete_docx_approval(document, approval_edit["values"])
        edits.append(
            {
                "kind": "complete_approval_table",
                "values": approval_edit["values"],
                "evidence": approval_edit["evidence"],
            }
        )

    # This source has an accidental duplicate code in a standalone cover
    # paragraph above the identity table.  Keep the table value and remove only
    # the stray paragraph.
    if record["target"]["code"] == "2001445-2":
        removed = 0
        for paragraph in document.paragraphs[:20]:
            if "2001445" not in paragraph.text:
                continue
            for run_item in paragraph.runs:
                run_item.text = ""
            removed += 1
        if removed != 1:
            raise RuntimeError(f"Expected one stray 2001445 cover paragraph, found {removed}")
        edits.append({"kind": "remove_stray_cover_code", "value": "2-2001445"})

    docx_path = directory / f"{record['target']['code']}.docx"
    document.save(docx_path)
    convert_dir = directory / f"convert-{record['target']['code']}"
    convert_dir.mkdir(parents=True, exist_ok=True)
    run(
        [
            str(SOFFICE),
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(convert_dir),
            str(docx_path),
        ],
        capture=True,
    )
    converted = convert_dir / f"{docx_path.stem}.pdf"
    if not converted.exists():
        raise RuntimeError(f"LibreOffice did not create {converted}")
    shutil.copy2(converted, output)
    return edits


def edit_source_pdf(record: dict[str, Any], output: Path, directory: Path) -> list[dict[str, Any]]:
    source = Path(record["source"]["absolute_path"])
    extracted = directory / f"{record['target']['code']}-extracted.pdf"
    extract_pdf_pages(source, record["source"].get("source_pages_inclusive") or "", extracted)
    document = pymupdf.open(extracted)
    edits: list[dict[str, Any]] = []
    title_edit = required_edit(record, "replace_cover_title")
    if title_edit:
        edit = replace_cover_value(document[0], "title", title_edit["to_value"])
        edit.update({"from": title_edit["from_value"], "to": title_edit["to_value"]})
        edits.append(edit)
    credit_edit = required_edit(record, "replace_credit_hours")
    if credit_edit:
        edits.append(
            replace_credit_hours(
                document,
                int(credit_edit["page_in_output"]),
                int(credit_edit["from_value"]),
                int(credit_edit["to_value"]),
            )
        )
    intermediate = directory / f"{record['target']['code']}-edited.pdf"
    document.save(intermediate, garbage=4, deflate=True, clean=True)
    document.close()
    shutil.copy2(intermediate, output)

    approval_edit = required_edit(record, "complete_approval_table")
    if approval_edit:
        page_number = int(record["approval"]["output_page_1_based"])
        changes = complete_pdf_approval(output, page_number, approval_edit["values"])
        edits.append(
            {
                "kind": "complete_approval_table",
                "page_1_based": page_number,
                "values": approval_edit["values"],
                "evidence": approval_edit["evidence"],
                "cells": changes,
            }
        )
    if required_edit(record, "extract_pages"):
        edits.insert(
            0,
            {
                "kind": "extract_pages",
                "source_pages_inclusive": record["source"]["source_pages_inclusive"],
            },
        )
    return edits


def resolve_repo_path(value: str, root: Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else root / path


def extract_arabic_stamp(manifest: dict[str, Any], output: Path, root: Path) -> Path:
    asset = manifest["stamp_assets"]["arabic_language"]
    document = pymupdf.open(resolve_repo_path(asset["path"], root))
    page = document[int(asset["page_1_based"]) - 1]
    width, height = asset["dimensions_pixels"]
    for image in page.get_images(full=True):
        extracted = document.extract_image(image[0])
        if extracted["width"] == width and extracted["height"] == height:
            target = output.with_suffix("." + extracted["ext"].replace("jpeg", "jpg"))
            target.write_bytes(extracted["image"])
            document.close()
            return target
    document.close()
    raise RuntimeError("Embedded Arabic Language Department stamp not found")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--output", type=Path, default=OUTPUT_RELATIVE)
    parser.add_argument("--working", type=Path, default=WORKING_RELATIVE)
    args = parser.parse_args()

    root = args.root.resolve()
    output = (root / args.output).resolve()
    working = (root / args.working).resolve()
    manifest_path = args.manifest
    if manifest_path is None:
        published_manifest = output / "source-manifest.json"
        manifest_path = published_manifest if published_manifest.exists() else TEMP_MANIFEST
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    records = manifest["records"]
    if len(records) != 55:
        raise RuntimeError(f"Expected 55 missing-course records, got {len(records)}")
    if any(record["target"]["current_course_details_present"] for record in records):
        raise RuntimeError("Manifest contains a course that already has site course_details")

    output.mkdir(parents=True, exist_ok=True)
    working.mkdir(parents=True, exist_ok=True)
    baseline = working / "baseline"
    baseline.mkdir(parents=True, exist_ok=True)
    source_docx = working / "docx"
    source_docx.mkdir(parents=True, exist_ok=True)
    arabic_stamp = extract_arabic_stamp(manifest, working / "arabic-language-stamp", root)

    build_records: list[dict[str, Any]] = []
    stamp_inventory: list[dict[str, Any]] = []
    for index, record in enumerate(records, start=1):
        code = record["target"]["code"]
        source = resolve_repo_path(record["source"]["absolute_path"], root)
        if sha256(source) != record["source"]["sha256"]:
            raise RuntimeError(f"Source hash changed: {source}")
        baseline_pdf = baseline / f"{code}.pdf"
        if record["source"]["format"] == "pdf":
            edits = edit_source_pdf(record, baseline_pdf, working)
        else:
            edits = convert_docx(record, baseline_pdf, source_docx)
        edits.extend(remove_nonpublic_link_annotations(baseline_pdf))
        run(["qpdf", "--check", str(baseline_pdf)], capture=True)
        reader = PdfReader(str(baseline_pdf), strict=False)
        baseline_pages = len(reader.pages)
        stamp = record["stamp"]
        if stamp["status"] != "preserve_existing_correct_stamp":
            stamp_asset = (
                arabic_stamp
                if stamp["department_key"] == "arabic_language"
                else resolve_repo_path(stamp["asset"]["path"], root)
            )
            target_page = record["approval"].get("output_page_1_based") or baseline_pages
            stamp_inventory.append(
                {
                    "relative_path": f"{code}.pdf",
                    "absolute_path": str(baseline_pdf),
                    "stamp_status": "missing_stamp",
                    "department_key": stamp["department_key"],
                    "expected_stamp_asset": str(stamp_asset),
                    "stamp_target_page_1_based": int(target_page),
                    "approval_location": "final_complete",
                    "baseline_sha256": sha256(baseline_pdf),
                    "page_count": baseline_pages,
                }
            )
        build_records.append(
            {
                "code": code,
                "title": record["target"]["title"],
                "hours": record["target"]["hours"],
                "match_type": record["match"]["type"],
                "source_relative_path": record["source"]["raw_relative_path"],
                "source_sha256": record["source"]["sha256"],
                "source_pages_inclusive": record["source"].get("source_pages_inclusive"),
                "approval": record["approval"],
                "edits_applied": edits,
                "stamp_status": stamp["status"],
                "baseline_sha256": sha256(baseline_pdf),
                "page_count_before_stamp": baseline_pages,
            }
        )
        print(f"[{index:02d}/{len(records)}] prepared {code}", flush=True)

    inventory_path = working / "stamp-inventory.json"
    inventory_path.write_text(
        json.dumps({"records": stamp_inventory}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    stamped = working / "stamped"
    stamp_audit = output / "stamp-audit.json"
    run(
        [
            sys.executable,
            str(root / "scripts/apply_department_stamps.py"),
            "--inventory",
            str(inventory_path),
            "--source-root",
            str(baseline),
            "--output-root",
            str(stamped),
            "--audit-output",
            str(stamp_audit),
        ]
    )
    stamp_data = json.loads(stamp_audit.read_text(encoding="utf-8"))
    stamp_by_code = {
        Path(item["relative_path"]).stem: item for item in stamp_data["records"]
    }

    for record in build_records:
        code = record["code"]
        source_pdf = stamped / f"{code}.pdf" if code in stamp_by_code else baseline / f"{code}.pdf"
        target_pdf = output / f"{code}.pdf"
        shutil.copy2(source_pdf, target_pdf)
        run(["qpdf", "--check", str(target_pdf)], capture=True)
        record["department_stamp"] = stamp_by_code.get(
            code, {"status": "preserved_existing_correct_stamp"}
        )
        record["output_sha256"] = sha256(target_pdf)
        record["page_count_after_stamp"] = len(PdfReader(str(target_pdf), strict=False).pages)

    source_manifest = output / "source-manifest.json"
    source_manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    result = {
        "schema": "raw-missing-course-recovery-bundle-v1",
        "selection_rule": manifest["selection_rule"],
        "records": build_records,
        "excluded_exact_code_conflicts": manifest["excluded_exact_code_conflicts"],
        "counts": {
            "total": len(build_records),
            "exact_identity": sum(
                item["match_type"] == "exact_code_and_identity" for item in build_records
            ),
            "approved_code_changes": sum(
                item["match_type"] == "approved_code_change_same_identity"
                for item in build_records
            ),
            "stamps_added": len(stamp_inventory),
            "stamps_preserved": len(build_records) - len(stamp_inventory),
            "approval_tables_completed": sum(
                any(edit["kind"] == "complete_approval_table" for edit in item["edits_applied"])
                for item in build_records
            ),
            "credit_hours_corrected": sum(
                any(edit["kind"] == "replace_credit_hours" for edit in item["edits_applied"])
                for item in build_records
            ),
        },
    }
    (output / "manifest.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    run(
        [
            sys.executable,
            str(root / "scripts/sanitize_audit_paths.py"),
            str(output),
        ]
    )
    print(json.dumps(result["counts"], ensure_ascii=False))


if __name__ == "__main__":
    main()
