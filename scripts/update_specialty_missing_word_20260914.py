#!/usr/bin/env python3
"""Update the formatted Word missing-course inventory after publication."""

from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "قائمة_المقررات_التخصصية_المفقودة_2026-09-14.docx"
RECOVERED_CODES = {"2002228-2", "2002229-2", "2002236-2", "20041201-2"}


def replace_run_text(paragraph, old: str, new: str) -> None:
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return
    raise ValueError(f"text not found in paragraph: {old!r}")


def set_cell_text(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def main() -> None:
    document = Document(TARGET)
    replace_run_text(document.paragraphs[3], "635 موضعًا متاحًا", "639 موضعًا متاحًا")
    replace_run_text(document.paragraphs[3], "121 موضعًا مفقودًا", "117 موضعًا مفقودًا")
    replace_run_text(document.paragraphs[3], "100 هوية مقرر", "96 هوية مقرر")
    replace_run_text(document.paragraphs[5], "83.99", "84.52")
    replace_run_text(document.paragraphs[8], "100 هوية", "96 هوية")
    replace_run_text(document.paragraphs[8], "121 موضعًا", "117 موضعًا")

    summary = document.tables[0]
    islamic_row = next(
        row
        for row in summary.rows[1:]
        if row.cells[0].text == "الدراسات الإسلامية"
        and row.cells[1].text == "بكالوريوس"
    )
    for cell, value in zip(islamic_row.cells[2:], ["102/127", "80.31%", "25"]):
        set_cell_text(cell, value)

    missing = document.tables[1]
    removed: set[str] = set()
    for row in list(missing.rows[1:]):
        code = row.cells[1].text.strip()
        if code in RECOVERED_CODES:
            removed.add(code)
            missing._tbl.remove(row._tr)
    if removed != RECOVERED_CODES:
        raise ValueError(f"missing Word rows not found: {sorted(RECOVERED_CODES - removed)}")
    if len(missing.rows) != 97:
        raise ValueError(f"expected 96 data rows, found {len(missing.rows) - 1}")
    for index, row in enumerate(missing.rows[1:], start=1):
        set_cell_text(row.cells[0], str(index))

    document.core_properties.modified = document.core_properties.created
    document.save(TARGET)
    print(f"updated {TARGET.name}: rows={len(missing.rows)-1}, coverage=639/756")


if __name__ == "__main__":
    main()
